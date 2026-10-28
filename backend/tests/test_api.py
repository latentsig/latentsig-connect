import io
import os
import sqlite3
import unittest
from pathlib import Path
from contextlib import closing
from unittest.mock import patch
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.base import Base
from app.db.models import Event, IntegrationJob, Interaction, Lead, User
from app.db.session import get_db
from app.core.config import Settings, settings
from app.core import rate_limit as rate_limit_module
from app.core.passwords import hash_password
from app.main import app
from app.services.integrations import worker
from app.services.backup import create_backup


class LeadApiTests(unittest.TestCase):
    def setUp(self):
        environment_patch = patch.dict(os.environ, {
            "DATABASE_URL": "sqlite:///./latentsig.db",
            "REDIS_URL": "",
        })
        self.addCleanup(environment_patch.stop)
        environment_patch.start()
        redis_patch = patch.object(settings, "redis_url", "")
        self.addCleanup(redis_patch.stop)
        redis_patch.start()
        # The in-memory limiter is process-wide; reset it between isolated API cases.
        rate_limit_module._hits.clear()
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.TestSession = sessionmaker(bind=self.engine, autoflush=False, autocommit=False)
        session_patch = patch("app.main.SessionLocal", self.TestSession)
        self.addCleanup(session_patch.stop)
        session_patch.start()

        with self.TestSession() as db:
            db.add_all([
                Event(name="Test Event", slug="test-event", location="Test City", is_active=True),
                User(email="bharat@latentsig.com", full_name="Admin User", role="ADMIN"),
                User(email="bhoomika@latentsig.com", full_name="Sales User", role="SALES"),
                User(email="raksha@latentsig.com", full_name="Read Only User", role="VIEWER"),
                User(email="other-sales@latentsig.com", full_name="Other Sales", role="SALES"),
            ])
            db.commit()

        def override_get_db():
            with self.TestSession() as db:
                yield db

        app.dependency_overrides[get_db] = override_get_db
        self.client_context = TestClient(app)
        self.client = self.client_context.__enter__()

    def tearDown(self):
        self.client_context.__exit__(None, None, None)
        app.dependency_overrides.clear()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()

    @staticmethod
    def payload(**overrides):
        value = {
            "full_name": "Taylor Example",
            "company": "Example Company",
            "job_title": "Product Manager",
            "email": "taylor@example.com",
            "email_type": "work",
            "phone": "+14155552671",
            "website": "example.com",
            "address": "One Example Street, Bengaluru",
            "data_processing_consent": True,
            "follow_up_consent": True,
            "interest": ["AI Agents"],
            "intent": "I need a demo",
            "event_slug": "test-event",
            "source": "api-test",
        }
        value.update(overrides)
        return value

    def submit(self, **overrides):
        return self.client.post("/api/v1/leads", json=self.payload(**overrides))

    def sign_in(self, email):
        return self.client.post("/api/v1/auth/email-login", json={"email": email})

    def test_submission_creates_normalized_lead_interaction_and_audit(self):
        response = self.submit(email="Taylor@Example.com", website="example.com", utm_source="qr", utm_medium="booth", utm_campaign="test-event")
        self.assertEqual(response.status_code, 201, response.text)
        self.assertTrue(response.json()["created"])

        with self.TestSession() as db:
            lead = db.scalar(select(Lead).where(Lead.normalized_email == "taylor@example.com"))
            self.assertIsNotNone(lead)
            self.assertEqual(lead.website, "https://example.com")
            self.assertEqual(lead.address, "One Example Street, Bengaluru")
            interaction = db.scalar(select(Interaction).where(Interaction.lead_id == lead.id))
            self.assertIsNotNone(interaction.data_processing_consent_at)
            self.assertIsNotNone(interaction.follow_up_consent_at)
            self.assertEqual(interaction.consent_notice_version, settings.consent_notice_version)
            self.assertEqual(interaction.utm_source, "qr")
            self.assertEqual(interaction.utm_medium, "booth")
            self.assertEqual(interaction.utm_campaign, "test-event")
            self.assertEqual(db.scalar(select(func.count()).select_from(Interaction).where(Interaction.lead_id == lead.id)), 1)

    def test_duplicate_contact_adds_interaction_without_creating_lead(self):
        first = self.submit()
        second = self.submit(
            full_name="Taylor Updated", email="TAYLOR@example.com", phone="+14155552672",
            event_slug="test-event", source="second-visit",
        )
        self.assertEqual(first.status_code, 201, first.text)
        self.assertEqual(second.status_code, 201, second.text)
        self.assertTrue(first.json()["created"])
        self.assertFalse(second.json()["created"])
        self.assertEqual(first.json()["id"], second.json()["id"])

        with self.TestSession() as db:
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), 1)
            self.assertEqual(db.scalar(select(func.count()).select_from(Interaction)), 2)

    def test_invalid_phone_and_interest_are_rejected(self):
        invalid_phone = self.submit(phone="not-a-phone")
        invalid_interest = self.submit(email="other@example.com", phone="+14155552672", interest=["Made up"])
        self.assertEqual(invalid_phone.status_code, 422)
        self.assertEqual(invalid_interest.status_code, 422)

        with self.TestSession() as db:
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), 0)

    def test_idempotency_key_returns_original_interaction_without_duplicate(self):
        key = str(uuid.uuid4())
        body = self.payload(email="idempotent@example.com", phone="+14155552672")
        first = self.client.post("/api/v1/leads", json=body, headers={"Idempotency-Key": key})
        retry = self.client.post("/api/v1/leads", json=body, headers={"Idempotency-Key": key})
        self.assertEqual(first.status_code, 201, first.text)
        self.assertEqual(retry.status_code, 201, retry.text)
        self.assertEqual(first.json(), retry.json())
        with self.TestSession() as db:
            self.assertEqual(db.scalar(select(func.count()).select_from(Interaction)), 1)

    def test_idempotency_key_must_be_a_uuid(self):
        response = self.client.post("/api/v1/leads", json=self.payload(), headers={"Idempotency-Key": "invalid"})
        self.assertEqual(response.status_code, 400)

    def test_other_interest_requires_details_and_persists_them(self):
        missing_details = self.submit(
            email="other-required@example.com", phone="+14155552673",
            interest=["Other"], other_interest_details="   ",
        )
        self.assertEqual(missing_details.status_code, 422)

        message = "RAG agent workflows"
        accepted = self.submit(
            email="other-details@example.com", phone="+14155552674",
            interest=["AI Agents", "Other"], other_interest_details=message,
            requirement_message="How the form felt and what could be improved.",
        )
        self.assertEqual(accepted.status_code, 201, accepted.text)
        with self.TestSession() as db:
            interaction = db.scalar(select(Interaction).where(Interaction.id == uuid.UUID(accepted.json()["interaction_id"])))
            self.assertEqual(interaction.interest, ["AI Agents", "Other"])
            self.assertEqual(interaction.other_interest_details, message)
            self.assertEqual(interaction.requirement_message, "How the form felt and what could be improved.")

        missing_discussion = self.submit(
            email="other-discussion-required@example.com", phone="+14155552675",
            intent="Other", requirement_message=" ",
        )
        too_many_words = self.submit(
            email="other-too-many@example.com", phone="+14155552676",
            interest=["Other"], other_interest_details="one two three four five six",
        )
        accepted_discussion = self.submit(
            email="other-discussion@example.com", phone="+14155552677",
            intent="Other", other_intent_details="RAG agent for support",
        )
        self.assertEqual(missing_discussion.status_code, 422)
        self.assertEqual(too_many_words.status_code, 422)
        self.assertEqual(accepted_discussion.status_code, 201, accepted_discussion.text)

    def test_business_card_submission_requires_tokens_and_manual_flow_rejects_them(self):
        missing_front = self.submit(capture_method="business_card")
        manual_with_card = self.submit(email="other@example.com", phone="+14155552672", business_card_token="not-a-token")
        self.assertEqual(missing_front.status_code, 422)
        self.assertEqual(manual_with_card.status_code, 422)

    def test_lead_submission_requires_explicit_consent(self):
        response = self.client.post("/api/v1/leads", json=self.payload(data_processing_consent=False))
        self.assertEqual(response.status_code, 422)

    def test_follow_up_consent_is_optional_and_recorded_separately(self):
        response = self.submit(email="no-follow-up@example.com", phone="+14155552678", follow_up_consent=False)
        self.assertEqual(response.status_code, 201, response.text)
        with self.TestSession() as db:
            interaction = db.scalar(select(Interaction).where(Interaction.id == uuid.UUID(response.json()["interaction_id"])))
            self.assertIsNotNone(interaction.data_processing_consent_at)
            self.assertIsNone(interaction.follow_up_consent_at)

    def test_card_scan_requires_explicit_ai_processing_consent(self):
        image = io.BytesIO()
        Image.new("RGB", (100, 100), "white").save(image, format="JPEG")
        response = self.client.post("/api/v1/ocr/business-card", files={"files": ("card.jpg", image.getvalue(), "image/jpeg")})
        self.assertEqual(response.status_code, 422)

    def test_ocr_upload_rejects_unsupported_and_mismatched_files_before_recognition(self):
        valid_jpeg = io.BytesIO()
        Image.new("RGB", (100, 100), "white").save(valid_jpeg, format="JPEG")
        unsupported = self.client.post("/api/v1/ocr/business-card", files={"files": ("card.txt", b"not an image", "text/plain")}, data={"ai_processing_consent": "true"})
        mismatch = self.client.post("/api/v1/ocr/business-card", files={"files": ("card.png", valid_jpeg.getvalue(), "image/png")}, data={"ai_processing_consent": "true"})
        self.assertEqual(unsupported.status_code, 415)
        self.assertEqual(mismatch.status_code, 415)

    def test_sheets_worker_retries_failure_without_losing_the_lead(self):
        submission = self.submit()
        interaction_id = submission.json()["interaction_id"]
        with self.TestSession() as db:
            db.add(IntegrationJob(type="google_sheets_upsert", entity_id=uuid.UUID(interaction_id), status="pending"))
            db.commit()

        with patch.object(worker, "SessionLocal", self.TestSession), patch.object(worker, "_sync_interaction", side_effect=RuntimeError("Sheets unavailable")):
            self.assertTrue(worker.process_one())
            self.assertFalse(worker.process_one())

        with self.TestSession() as db:
            lead = db.scalar(select(Lead).where(Lead.email == "taylor@example.com"))
            job = db.scalar(select(IntegrationJob).where(IntegrationJob.entity_id == uuid.UUID(interaction_id)))
            self.assertIsNotNone(lead)
            self.assertEqual(job.status, "pending")
            self.assertEqual(job.attempt_count, 1)
            self.assertEqual(job.last_error, "Sheets unavailable")
            self.assertIsNotNone(job.next_retry_at)

    def test_sheets_worker_marks_successful_job_complete(self):
        submission = self.submit()
        interaction_id = submission.json()["interaction_id"]
        with self.TestSession() as db:
            db.add(IntegrationJob(type="google_sheets_upsert", entity_id=uuid.UUID(interaction_id), status="pending"))
            db.commit()

        with patch.object(worker, "SessionLocal", self.TestSession), patch.object(worker, "_sync_interaction"):
            self.assertTrue(worker.process_one())

        with self.TestSession() as db:
            job = db.scalar(select(IntegrationJob).where(IntegrationJob.entity_id == uuid.UUID(interaction_id)))
            self.assertEqual(job.status, "succeeded")
            self.assertEqual(job.attempt_count, 1)

    def test_lead_sort_dashboard_breakdowns_and_bulk_assignment(self):
        first = self.submit(full_name="Zoe Example", email="zoe@example.com", phone="+14155552671")
        second = self.submit(full_name="Aaron Example", email="aaron@example.com", phone="+14155552672")
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 201)
        self.assertEqual(self.sign_in("bharat@latentsig.com").status_code, 200)

        sorted_leads = self.client.get("/api/v1/admin/leads", params={"sort_by": "full_name", "sort_order": "asc"})
        self.assertEqual(sorted_leads.status_code, 200, sorted_leads.text)
        self.assertEqual([row["full_name"] for row in sorted_leads.json()["items"]], ["Aaron Example", "Zoe Example"])
        dashboard = self.client.get("/api/v1/admin/dashboard")
        self.assertEqual(dashboard.json()["by_event"], {"Test Event": 2})
        self.assertEqual(dashboard.json()["by_status"]["new"], 2)

        manager = None
        with self.TestSession() as db:
            manager = User(email="manager@latentsig.com", full_name="Manager", role="MANAGER")
            db.add(manager)
            db.commit()
            db.refresh(manager)
            manager_id = str(manager.id)
        assigned = self.client.post("/api/v1/admin/leads/bulk-assign", json={"lead_ids": [first.json()["id"], second.json()["id"]], "assigned_to": manager_id})
        self.assertEqual(assigned.status_code, 200, assigned.text)
        self.assertEqual(assigned.json()["updated"], 2)

        with self.TestSession() as db:
            rows = db.scalars(select(Interaction)).all()
            self.assertEqual({str(row.assigned_to) for row in rows}, {manager_id})

    def test_sales_export_contains_all_visible_rows_and_excludes_other_assignees(self):
        shared = self.submit(email="shared@example.com", phone="+14155552671")
        assigned_elsewhere = self.submit(email="private@example.com", phone="+14155552672")
        self.assertEqual(shared.status_code, 201)
        self.assertEqual(assigned_elsewhere.status_code, 201)

        with self.TestSession() as db:
            other_sales = db.scalar(select(User).where(User.email == "other-sales@latentsig.com"))
            hidden_lead = db.scalar(select(Lead).where(Lead.email == "private@example.com"))
            hidden_interaction = db.scalar(select(Interaction).where(Interaction.lead_id == hidden_lead.id))
            hidden_interaction.assigned_to = other_sales.id
            db.commit()

        self.assertEqual(self.sign_in("bhoomika@latentsig.com").status_code, 200)
        response = self.client.get("/api/v1/admin/exports/leads.xlsx")
        self.assertEqual(response.status_code, 200, response.text)
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Leads", "Summary", "Export Metadata"])
        lead_sheet = workbook["Leads"]
        self.assertEqual(lead_sheet.max_row, 2)
        self.assertEqual(lead_sheet.cell(2, 6).value, "shared@example.com")
        self.assertEqual(lead_sheet.cell(1, 20).value, "Address")
        self.assertEqual(lead_sheet.cell(2, 20).value, "One Example Street, Bengaluru")
        workbook.close()

    def test_viewer_cannot_export_and_unallowlisted_user_cannot_sign_in(self):
        self.assertEqual(self.sign_in("raksha@latentsig.com").status_code, 200)
        denied = self.client.get("/api/v1/admin/exports/leads.xlsx")
        self.assertEqual(denied.status_code, 403)

        self.client.post("/api/v1/auth/logout")
        rejected = self.sign_in("outside@example.com")
        self.assertEqual(rejected.status_code, 403)

    def test_email_login_accepts_exact_team_allowlist_only_in_development(self):
        from app.core.config import settings

        expected_emails = {
            "raksha@latentsig.com", "bhoomika@latentsig.com", "bharat@latentsig.com",
            "info@latentsig.com", "shauryasharthak@latentsig.com",
        }
        with patch.object(settings, "admin_email_allowlist", ",".join(sorted(expected_emails))):
            for email in sorted(expected_emails):
                response = self.client.post("/api/v1/auth/email-login", json={"email": email})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["email"], email)

            outsider = self.client.post("/api/v1/auth/email-login", json={"email": "outside@example.com"})
            self.assertEqual(outsider.status_code, 403)
            self.assertEqual(self.client.get("/api/v1/auth/google/start").status_code, 404)
            self.assertTrue(self.client.get("/api/v1/auth/providers").json()["development_email_login"])
            self.assertEqual(self.client.post("/api/v1/auth/password-login", json={
                "email": "bharat@latentsig.com", "password": "incorrect password",
            }).status_code, 401)

        with patch.object(settings, "app_env", "production"), patch.object(
            rate_limit_module, "settings", Settings(_env_file=None, app_env="development"),
        ):
            disabled = self.client.post("/api/v1/auth/email-login", json={"email": "bharat@latentsig.com"})
            self.assertEqual(disabled.status_code, 404)

    def test_health_readiness_and_http_metrics(self):
        self.assertEqual(self.client.get("/health/live").status_code, 200)
        self.assertEqual(self.client.get("/health/ready").json(), {"status": "ready"})
        response = self.client.get("/metrics")
        self.assertEqual(response.status_code, 200)
        self.assertIn('route="/health/ready"', response.text)

    def test_auth_providers_expose_development_email_login_only_in_development(self):
        self.assertEqual(self.client.get("/api/v1/auth/providers").json(), {
            "password_login": True,
            "development_email_login": True,
        })
        with patch.object(settings, "app_env", "production"), patch.object(
            rate_limit_module, "settings", Settings(_env_file=None, app_env="development"),
        ):
            self.assertEqual(self.client.get("/api/v1/auth/providers").json(), {
                "password_login": True,
                "development_email_login": False,
            })

    def test_production_password_login_sets_session_for_active_allowlisted_user(self):
        password = "correct horse battery staple"
        with self.TestSession() as db:
            user = db.scalar(select(User).where(User.email == "bharat@latentsig.com"))
            user.password_hash = hash_password(password)
            db.commit()

        with patch.object(settings, "app_env", "production"), patch.object(
            rate_limit_module, "settings", Settings(_env_file=None, app_env="development"),
        ):
            response = self.client.post("/api/v1/auth/password-login", json={
                "email": "bharat@latentsig.com", "password": password,
            })
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["email"], "bharat@latentsig.com")
            self.assertEqual(self.client.get("/api/v1/auth/me").status_code, 200)

            wrong_password = self.client.post("/api/v1/auth/password-login", json={
                "email": "bharat@latentsig.com", "password": "wrong password long",
            })
            outside_allowlist = self.client.post("/api/v1/auth/password-login", json={
                "email": "outside@example.com", "password": password,
            })
            self.assertEqual(wrong_password.status_code, 401)
            self.assertEqual(outside_allowlist.status_code, 401)

    def test_sqlite_backup_is_integrity_checked(self):
        from app.core.config import settings

        root = Path(__file__).resolve().parent
        source = root / f".backup-source-{uuid.uuid4().hex}.db"
        backup = None
        try:
            with closing(sqlite3.connect(source)) as database:
                database.execute("CREATE TABLE sample (value TEXT)")
                database.execute("INSERT INTO sample VALUES ('preserved')")
                database.commit()
            with patch.object(settings, "database_url", f"sqlite:///{source.as_posix()}"):
                backup = create_backup(root)
            self.assertTrue(backup.is_file())
            with closing(sqlite3.connect(backup)) as database:
                self.assertEqual(database.execute("PRAGMA integrity_check").fetchone()[0], "ok")
                self.assertEqual(database.execute("SELECT value FROM sample").fetchone()[0], "preserved")
        finally:
            for temporary_path in (source, backup, backup.with_suffix(backup.suffix + ".partial") if backup else None):
                if temporary_path and temporary_path.exists():
                    temporary_path.unlink()

    def test_production_settings_do_not_require_google_signup_or_sheets(self):
        with self.assertRaisesRegex(ValueError, "Production requires PostgreSQL"):
            Settings(_env_file=None, app_env="production", app_secret="s" * 40)
        settings = Settings(
            _env_file=None,
            app_env="production", app_secret="s" * 40,
            database_url=f"postgresql+psycopg://latentsig:{'s' * 40}@db:5432/latentsig",
            redis_url="redis://redis:6379/0", frontend_url="https://connect.example.com",
            openrouter_api_key="test-only-key",
            frontend_urls="https://connect.example.com",
            admin_email_allowlist="admin@example.com", initial_admin_emails="admin@example.com",
        )
        self.assertEqual(settings.app_env, "production")

    def test_production_settings_allow_manual_registration_without_openrouter(self):
        production = Settings(
            _env_file=None,
            app_env="production",
            app_secret="s" * 40,
            database_url=f"postgresql+psycopg://latentsig:{'s' * 40}@db:5432/latentsig",
            redis_url="redis://redis:6379/0",
            frontend_url="https://connect.latentsig.com",
            frontend_urls="https://connect.latentsig.com",
            admin_email_allowlist="bhoomika@latentsig.com,bharat@latentsig.com",
            initial_admin_emails="bhoomika@latentsig.com",
            ocr_enabled=False,
            openrouter_api_key="",
        )
        self.assertEqual(production.admin_emails, {"bhoomika@latentsig.com", "bharat@latentsig.com"})
        production_values = {
            "_env_file": None,
            "app_env": "production",
            "app_secret": "s" * 40,
            "database_url": f"postgresql+psycopg://latentsig:{'s' * 40}@db:5432/latentsig",
            "redis_url": "redis://redis:6379/0",
            "frontend_url": "https://connect.latentsig.com",
            "frontend_urls": "https://connect.latentsig.com",
            "admin_email_allowlist": "bhoomika@latentsig.com,bharat@latentsig.com",
            "initial_admin_emails": "bhoomika@latentsig.com",
            "ocr_enabled": False,
        }
        with self.assertRaisesRegex(ValueError, "APP_SECRET must be a unique"):
            Settings(**{**production_values, "app_secret": "replace-with-generated-random-value"})
        with self.assertRaisesRegex(ValueError, "unique PostgreSQL password"):
            Settings(**{**production_values, "database_url": "postgresql+psycopg://latentsig:replace-with-generated-random-value@db:5432/latentsig"})
        with self.assertRaisesRegex(ValueError, "OPENROUTER_API_KEY is required"):
            Settings(
                _env_file=None,
                app_env="production",
                app_secret="s" * 40,
                database_url=f"postgresql+psycopg://latentsig:{'s' * 40}@db:5432/latentsig",
                redis_url="redis://redis:6379/0",
                frontend_url="https://connect.latentsig.com",
                frontend_urls="https://connect.latentsig.com",
                admin_email_allowlist="bhoomika@latentsig.com,bharat@latentsig.com",
                initial_admin_emails="bhoomika@latentsig.com",
                ocr_enabled=True,
                openrouter_api_key="",
            )


if __name__ == "__main__":
    unittest.main()
